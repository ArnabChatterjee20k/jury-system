"""Document text extraction + NLP entity extraction layer.

- Text: pdfplumber (PDF), python-docx (DOCX), plain text (TXT/MD).
- NLP: spaCy NER (if en_core_web_sm is installed) plus legal-specific regexes
  for IPC/BNS/CrPC sections, dates, times and FIR numbers.
"""
from __future__ import annotations

import logging
import re
from collections import Counter
from pathlib import Path

from .errors import DocumentError

log = logging.getLogger(__name__)

TEXT_EXTS = {".txt", ".md", ".text"}
SUPPORTED_EXTS = {".pdf", ".docx", *TEXT_EXTS}
# Evidence exhibits (images/audio/video) are accepted but not text-extracted.
EXHIBIT_EXTS = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".mp3", ".wav", ".mp4", ".mov"}


def extract_text(path: Path) -> tuple[str, str | None]:
    """Return (text, warning)."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import pdfplumber

            pages = []
            with pdfplumber.open(path) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    t = page.extract_text() or ""
                    if t.strip():
                        pages.append(f"[Page {i}]\n{t}")
            text = "\n\n".join(pages)
            if not text.strip():
                return "", "No extractable text found — the PDF may be scanned. OCR it first for full analysis."
            return text, None
        if ext == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts), None
        if ext in TEXT_EXTS:
            return path.read_text(encoding="utf-8", errors="replace"), None
        if ext in EXHIBIT_EXTS:
            return "", "Media exhibit recorded by filename only; content is not analysed."
    except Exception as exc:  # corrupt / password-protected files etc.
        log.exception("Failed to extract %s", path)
        raise DocumentError(
            f"Could not read '{path.name}': {exc}",
            hint="Check that the file isn't corrupted or password-protected.",
        ) from exc
    raise DocumentError(
        f"Unsupported file type '{ext}' for '{path.name}'.",
        hint=f"Supported: {', '.join(sorted(SUPPORTED_EXTS | EXHIBIT_EXTS))}",
    )


# ---------------------------------------------------------------- NLP layer

_nlp = None
_nlp_error: str | None = None


def get_nlp():
    global _nlp, _nlp_error
    if _nlp is None and _nlp_error is None:
        try:
            import spacy

            _nlp = spacy.load("en_core_web_sm")
        except Exception as exc:
            _nlp_error = f"spaCy model unavailable ({exc}); using regex-only extraction."
            log.warning(_nlp_error)
    return _nlp


def nlp_status() -> dict:
    get_nlp()
    return {"spacy": _nlp is not None, "message": _nlp_error}


SECTION_RE = re.compile(
    r"\b(?:u/s|under\s+sections?|sections?|sec\.?|s\.)\s*"
    r"(\d{1,3}[A-Z]?(?:\s*\(\d+\))?(?:\s*(?:,|/|and|&)\s*\d{1,3}[A-Z]?(?:\s*\(\d+\))?)*)"
    r"\s*(?:of\s+(?:the\s+)?)?"
    r"(IPC|BNS|BNSS|BSA|CrPC|Cr\.P\.C\.?|Indian Penal Code|Bharatiya Nyaya Sanhita|"
    r"Bharatiya Nagarik Suraksha Sanhita|Bharatiya Sakshya Adhiniyam|Evidence Act|NDPS Act|Arms Act|IT Act)?",
    re.IGNORECASE,
)
TIME_RE = re.compile(r"\b(?:[01]?\d|2[0-3])(?::[0-5]\d)?\s*(?:a\.?m\.?|p\.?m\.?|hrs|hours)\b|\b(?:[01]\d|2[0-3]):[0-5]\d\b", re.IGNORECASE)
DATE_RE = re.compile(
    r"\b\d{1,2}[./-]\d{1,2}[./-]\d{2,4}\b|"
    r"\b\d{1,2}(?:st|nd|rd|th)?\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*,?\s+\d{4}\b|"
    r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\s+\d{1,2}(?:st|nd|rd|th)?,?\s+\d{4}\b",
    re.IGNORECASE,
)
FIR_RE = re.compile(r"\bFIR\s*(?:No\.?|Number|#)?\s*[:\-]?\s*([A-Z0-9/\-]{2,30})", re.IGNORECASE)

_CODE_NORMALISE = {
    "indian penal code": "IPC",
    "bharatiya nyaya sanhita": "BNS",
    "bharatiya nagarik suraksha sanhita": "BNSS",
    "bharatiya sakshya adhiniyam": "BSA",
    "cr.p.c": "CrPC",
    "cr.p.c.": "CrPC",
    "crpc": "CrPC",
}


def _top(counter: Counter, n: int = 25) -> list[dict]:
    return [{"text": k, "count": v} for k, v in counter.most_common(n)]


def extract_entities(docs: list[dict]) -> dict:
    """docs: [{filename, category, text}] -> structured entity summary."""
    people, orgs, places, dates, times, sections, firs = (Counter() for _ in range(7))
    nlp = get_nlp()

    for d in docs:
        text = d["text"]
        if not text:
            continue
        for m in SECTION_RE.finditer(text):
            code = (m.group(2) or "").strip()
            code = _CODE_NORMALISE.get(code.lower(), code.upper() if len(code) <= 5 else code)
            for num in re.split(r"\s*(?:,|/|and|&)\s*", m.group(1)):
                num = num.strip()
                if num:
                    sections[f"Section {num}{' ' + code if code else ''}"] += 1
        for m in TIME_RE.finditer(text):
            times[m.group(0).strip()] += 1
        for m in DATE_RE.finditer(text):
            dates[m.group(0).strip()] += 1
        for m in FIR_RE.finditer(text):
            if any(ch.isdigit() for ch in m.group(1)):
                firs[m.group(1).strip("/-")] += 1

        if nlp is not None:
            # Chunk to stay well under spaCy's max_length and keep memory bounded.
            for i in range(0, min(len(text), 300_000), 50_000):
                for ent in nlp(text[i : i + 50_000]).ents:
                    val = " ".join(ent.text.split())
                    if len(val) < 2:
                        continue
                    if ent.label_ == "PERSON":
                        people[val] += 1
                    elif ent.label_ == "ORG":
                        orgs[val] += 1
                    elif ent.label_ in ("GPE", "LOC", "FAC"):
                        places[val] += 1
                    elif ent.label_ == "DATE":
                        dates[val] += 1
                    elif ent.label_ == "TIME":
                        times[val] += 1

    return {
        "engine": "spacy+regex" if nlp is not None else "regex",
        "people": _top(people),
        "organizations": _top(orgs),
        "places": _top(places),
        "dates": _top(dates),
        "times": _top(times),
        "sections": _top(sections, 40),
        "fir_numbers": _top(firs, 10),
    }
